"""RAG 解析层与分块层的表格处理回归用例。

覆盖三处改动：
1. docx 按 body 子节点顺序流式遍历（表格不再被搬到文末）；
2. splitter 含表格文本走行对齐切分（一行记录不被拦腰截断）；
3. pdf 接入 pdfplumber 抽表转 Markdown，并保证表格内容不被重复抽取。
"""
import logging
import sys
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.rag import loader, splitter

SAMPLES = Path(__file__).resolve().parent.parent / "samples" / "demo_docs" / "generated"
SALES_PDF = SAMPLES / "2025年各渠道销售与库存盘点表.pdf"
SALES_DOCX = SAMPLES / "2025年各渠道销售与库存盘点表.docx"

needs_sales_pdf = pytest.mark.skipif(not SALES_PDF.exists(), reason="缺少演示 PDF 样本")
needs_sales_docx = pytest.mark.skipif(not SALES_DOCX.exists(), reason="缺少演示 DOCX 样本")


# ------------------------------------------------------------------ 辅助构造


def _build_docx(path: Path, table_rows: list[list[str]]) -> None:
    """构造「段落 - 表格 - 段落」夹心结构的 docx。"""
    doc = DocxDocument()
    doc.add_paragraph("第一章 总则")
    doc.add_paragraph("本章说明适用范围。")
    table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
    for r, row in enumerate(table_rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_paragraph("第二章 附则")
    doc.save(str(path))


def _reference_split_by_chars(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    """改造前的字符滑窗实现（用于确认无表格文本行为未被改变）。"""
    text = (text or "").strip()
    if not text:
        return []
    if chunk_overlap >= chunk_size:
        chunk_overlap = 0
    chunks: list[str] = []
    step = max(1, chunk_size - chunk_overlap)
    start = 0
    while start < len(text):
        chunk = text[start:start + chunk_size].strip()
        if chunk:
            chunks.append(chunk)
        if start + chunk_size >= len(text):
            break
        start += step
    return chunks


# ------------------------------------------------------------------- 改进 1


def test_docx_table_keeps_position_between_paragraphs(tmp_path):
    """表格必须留在原位，而不是被两趟实现搬到文末。"""
    path = tmp_path / "sandwich.docx"
    _build_docx(path, [["渠道", "销售额"], ["天猫", "4820"], ["京东", "3960"]])

    text = loader.extract_text(str(path), "docx")

    i_head = text.index("第一章 总则")
    i_title = text.index("本章说明适用范围")
    i_table = text.index("渠道 | 销售额")
    i_tail = text.index("第二章 附则")
    assert i_head < i_title < i_table < i_tail
    assert "天猫 | 4820" in text
    assert "京东 | 3960" in text


def test_docx_merged_cell_output_once(tmp_path):
    """横向合并单元格指向同一 w:tc，不应重复输出成多列。"""
    path = tmp_path / "merged.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "合计口径"
    table.cell(0, 0).merge(table.cell(0, 2))
    for c, value in enumerate(["全年", "963", "9.9"]):
        table.cell(1, c).text = value
    doc.save(str(path))

    text = loader.extract_text(str(path), "docx")

    assert text.count("合计口径") == 1
    assert "全年 | 963 | 9.9" in text


@needs_sales_docx
def test_sales_docx_table_near_its_heading():
    """实测回归：旧实现表格漂到第 804 字符，新实现应紧跟章节标题。"""
    text = loader.extract_text(str(SALES_DOCX), "docx")
    i_heading = text.index("一、年度分渠道销售与库存明细")
    i_table = text.index("渠道 | Q1销售额")
    assert 0 < i_table - i_heading < 300, f"表格与标题距离 {i_table - i_heading} 字符，疑似仍被搬到文末"


# ------------------------------------------------------------------- 改进 2


def test_has_table_detection():
    assert splitter.has_table(["| a | b |", "| c | d |"])
    assert not splitter.has_table(["普通段落，没有竖线", "第二行也没有"])
    assert not splitter.has_table(["只有一个 a | b", "还剩一个 c | d"])  # 单竖线不算表格


def test_plain_text_split_unchanged():
    """无表格文本必须与改造前的字符滑窗完全一致。"""
    text = "数据库应急响应手册。" * 120
    assert splitter.split_text(text, 500, 50) == _reference_split_by_chars(text, 500, 50)


def test_table_rows_never_cut_in_half():
    """块内每一行都必须是完整记录，不允许出现半截记录（腰斩）。"""
    header = "| 渠道 | 销售额 | 退货率 | 周转 |"
    divider = "| --- | --- | --- | --- |"
    rows = [f"| 渠道{i} | {i * 1000} | {i * 7}% | {i * 3} |" for i in range(1, 25)]
    text = "\n".join([header, divider, *rows])
    valid = {header, divider, *rows}

    chunks = splitter.split_text(text, 300, 30)

    for idx, chunk in enumerate(chunks):
        for line in chunk.split("\n"):
            if "|" in line:
                assert line in valid, f"块 {idx} 出现半截记录: {line!r}"
    assert all(any(row in chunk for chunk in chunks) for row in rows)


def test_table_chunks_respect_size_limit():
    rows = [f"| 行{i} | " + "数据" * 20 + " |" for i in range(1, 40)]
    text = "\n".join(rows)
    chunks = splitter.split_text(text, 400, 40)
    for chunk in chunks:
        assert len(chunk) <= 400, f"块长 {len(chunk)} 超过 chunk_size，行对齐打包失效"


# --------------------------------------------------------- 表头回填（改进 2 续）


def _big_table(rows: int = 40):
    header = "| 渠道 | Q1销售额 | Q2销售额 | Q3销售额 | Q4销售额 | 全年合计 | 退货率 | 期末SKU库存 | 周转天数 |"
    divider = "| --- | --- | --- | --- | --- | --- | --- | --- | --- |"
    body = [
        f"| 渠道{i} | {i * 100} | {i * 110} | {i * 120} | {i * 130} | {i * 460} | {i}% | {i * 9} | {i} |"
        for i in range(1, rows + 1)
    ]
    return header, divider, body


def test_every_chunk_of_big_table_starts_with_header():
    """超过 chunk_size 的大表格：每个块的开头都必须是表头 + 分隔线。"""
    header, divider, body = _big_table(40)
    text = "\n".join([header, divider, *body])

    chunks = splitter.split_text(text, 400, 40)

    assert len(chunks) >= 4, f"样本未触发多次切分，块数={len(chunks)}"
    for k, chunk in enumerate(chunks):
        head_lines = chunk.split("\n")[:2]
        assert head_lines[0] == header, f"块{k} 首行不是表头: {head_lines[0][:40]!r}"
        assert head_lines[1] == divider, f"块{k} 第二行不是分隔线: {head_lines[1][:40]!r}"
        assert chunk.count(header) == 1, f"块{k} 表头重复插入"
    # 每条记录仍然完整、且都出现在某个块里
    for row in body:
        assert any(row in chunk for chunk in chunks), f"记录丢失: {row[:30]}"


def test_mid_table_chunk_gets_header_when_no_divider():
    """docx / xlsx 抽出的行没有分隔线，块首应只补表头一行。"""
    header = "渠道 | 销售额 | 退货率 | 周转天数"
    body = [f"渠道{i} | {i * 100} | {i}% | {i * 2}" for i in range(1, 30)]
    text = "\n".join([header, *body])

    chunks = splitter.split_text(text, 200, 20)

    assert len(chunks) >= 3
    for k, chunk in enumerate(chunks):
        assert chunk.split("\n")[0] == header, f"块{k} 缺少表头"
        assert chunk.count(header) == 1


def test_table_prefixes_map_is_exact():
    """前缀映射的三种情形：表头补 None、分隔线补表头、数据行补表头+分隔线。"""
    header = "| 渠道 | 金额 |"
    divider = "| --- | --- |"
    data = "| 天猫 | 4820 |"
    lines = ["前言", header, divider, data, "后记", "x | y | z", "u | v | w"]

    prefixes = splitter._table_prefixes(lines)

    assert prefixes[0] is None          # 表格外正文
    assert prefixes[1] is None          # 表头本身不补
    assert prefixes[2] == header        # 分隔线只补表头
    assert prefixes[3] == f"{header}\n{divider}"
    assert prefixes[4] is None          # 表格结束


def test_prose_not_touched_by_header_backfill():
    """不含表格的正文不应被补任何前缀。"""
    text = "第一段内容。\n第二段内容。\n第三段内容。"
    assert splitter._table_prefixes(text.split("\n")) == [None, None, None]
    assert all("|" not in c for c in splitter.split_text(text, 20, 5))


def test_long_single_line_still_split():
    """PDF 抽出的超长单行无法行对齐，必须退化为硬切而不是产出一个超大块。"""
    text = "| a | b |\n| --- | --- |\n" + "超长单行" * 500
    chunks = splitter.split_text(text, 500, 50)
    assert len(chunks) > 1
    assert all(len(c) <= 500 for c in chunks)


def test_overlap_does_not_duplicate_tail_chunk():
    """重叠尾行不应单独再产出一个与上一块重复的块。"""
    text = "\n".join(f"| 行{i} | 值{i} |" for i in range(50))
    chunks = splitter.split_text(text, 200, 50)
    assert chunks[-1] != chunks[-2]
    assert len(chunks) == len({*chunks})


# ------------------------------------------------------------------- 改进 3


def test_table_to_markdown_shape():
    md = loader._table_to_markdown([["渠道", "金额", None], ["天猫", "4820"], []])
    lines = md.split("\n")
    assert lines[0] == "| 渠道 | 金额 |  |"
    assert lines[1] == "| --- | --- | --- |"
    assert lines[2] == "| 天猫 | 4820 |  |"
    assert len(lines) == 3


def test_table_to_markdown_escapes_pipe_and_newline():
    md = loader._table_to_markdown([["a|b", "c\nd"]])
    assert "a\\|b" in md
    assert "c d" in md


def test_table_to_markdown_empty():
    assert loader._table_to_markdown([]) == ""
    assert loader._table_to_markdown([[None, ""]]) == ""


@needs_sales_pdf
def test_pdf_table_becomes_markdown():
    text = loader.extract_text(str(SALES_PDF), "pdf")
    assert "| 渠道 | Q1销售额 | Q2销售额 | Q3销售额 | Q4销售额 | 全年合计 | 退货率 | 期末SKU库存 | 周转天数 |" in text
    assert "| --- | --- | --- | --- | --- | --- | --- | --- | --- |" in text
    assert "| 天猫旗舰店 | 4,820 | 5,310 | 5,980 | 11,240 | 27,350 | 8.4% | 12,480 | 42 |" in text
    assert "| 合计 | 16,250 | 17,930 | 20,410 | 41,670 | 96,260 | 9.8% | 46,650 | 39.4 |" in text


@needs_sales_pdf
def test_pdf_table_text_not_duplicated():
    """表格内容只能以 Markdown 形式出现一次，不能同时残留空格分隔的纯文本版本。"""
    text = loader.extract_text(str(SALES_PDF), "pdf")
    assert text.count("天猫旗舰店") == 1
    assert "渠道 Q1销售额 Q2销售额" not in text


@needs_sales_pdf
def test_pdf_table_keeps_heading_order():
    text = loader.extract_text(str(SALES_PDF), "pdf")
    assert text.index("一、年度分渠道销售与库存明细") < text.index("| 渠道 |") < text.index("二、年度汇总指标")


@needs_sales_pdf
def test_pdf_logs_table_hint(caplog):
    with caplog.at_level(logging.INFO, logger="app.rag.loader"):
        loader.extract_text(str(SALES_PDF), "pdf")
    messages = [r.getMessage() for r in caplog.records]
    assert any("[解析提示]" in m and "表格" in m and "Markdown" in m for m in messages), messages


@needs_sales_pdf
def test_pdf_falls_back_to_pypdf_when_pdfplumber_missing(monkeypatch):
    """pdfplumber 缺失时必须退回 pypdf，而不是抛错。"""
    monkeypatch.setitem(sys.modules, "pdfplumber", None)

    text = loader.extract_text(str(SALES_PDF), "pdf")

    assert "天猫旗舰店" in text
    assert "渠道 Q1销售额 Q2销售额" in text  # pypdf 的空格分隔版本
    assert "| 渠道 |" not in text
